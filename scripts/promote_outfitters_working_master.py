#!/usr/bin/env python3
"""Promote the reviewed hunting-outfitter workbook into repository data.

The spreadsheet owns business identity, contact fields, review status, and
reported USFS/BLM permit areas. Existing repository records may contribute
non-contact enrichment only when their business identity resolves uniquely.
This promotion pass leaves unit associations blank. The separate spatial
eligibility builder maps confirmed federal permit areas to every intersecting
DWR species/unit geometry; whole-unit percentage is descriptive only.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
LOCAL_OUTFITTER_DIR = ROOT / "local_data" / "outfitters"
MASTER_PATH = LOCAL_OUTFITTER_DIR / "outfitters-master.internal.json"
LEGACY_MASTER_PATH = ROOT / "data" / "outfitters-master.json"
PUBLIC_PATH = ROOT / "data" / "outfitters-public.json"
EVIDENCE_PATH = LOCAL_OUTFITTER_DIR / "outfitter-federal-service-area-evidence.internal.json"
MANIFEST_PATH = ROOT / "data" / "outfitter-internal-master-manifest.json"
MANTI_LA_SAL_AUTHORITY_PATH = (
    ROOT / "data" / "source-evidence" / "manti-la-sal-permitted-hunting-outfitters-2026.json"
)
FISHLAKE_AUTHORITY_PATH = (
    ROOT / "data" / "source-evidence" / "fishlake-permitted-hunting-outfitters-2025.json"
)
WILD_EYEZ_CONFIRMED_SERVICE_PATH = (
    ROOT / "data" / "source-evidence" / "wild-eyez-confirmed-elk-service-area-2026.json"
)
COVERAGE_PATH = LOCAL_OUTFITTER_DIR / "outfitter-federal-unit-coverage-review.internal.json"
CANONICAL_PATH = ROOT / "canonical" / "outfitter-verification-2026.json"

REQUIRED_HEADERS = {
    "HUNTING OR FISHING",
    "OUTFITTER",
    "OWNER",
    "PHONE NUMBER",
    "WEBSITE",
    "EMAIL",
    "SUP’s USFS",
    "SRP's BLM",
    "LOCATION",
    "Business Entity",
    "REVIEW STATUS",
    "RECONCILIATION NOTES",
}
REVIEW_STATUSES = {"Confirmed", "Needs Verification", "Spreadsheet Only"}

# These are identity aliases only. Contact-field overlap is deliberately not
# used because several legacy repository records contain merged contacts.
LEGACY_NAME_ALIASES = {
    "battleground g/o": "BATTLEGROUND GUIDES & OUTFITTERS, LLC",
    "book cliff outfitters": "Book Cliff Outfitters/Cisco Outfitters",
    "bull mountain outfiiitters": "BULL MOUNTAIN OUTFITTERS, LLC",
    "tripple h": "Triple H Hunting",
}

USFS_FORESTS = (
    ("ashley", "Ashley", ("ashley", "duchesne", "duschene")),
    ("dixie", "Dixie", ("dixie",)),
    ("fishlake", "Fishlake", ("fishlake",)),
    ("manti-la-sal", "Manti-La Sal", ("manti", "la sal", "lasal")),
    (
        "uwc",
        "Uinta-Wasatch-Cache",
        (
            "uwc",
            "uinta-wasatch-cache",
            "uinta wasatch cache",
            "spanish fork",
            "pleasant grove",
            "heber",
            "kamas",
            "evanston",
        ),
    ),
)

USFS_DISTRICTS = (
    ("ashley-duchesne", "Duchesne", ("duchesne", "duschene")),
    ("ashley-roosevelt", "Roosevelt", ("roosevelt",)),
    ("ashley-vernal", "Vernal", ("vernal",)),
    ("dixie-cedar", "Cedar City", ("dixie", "cedar")),
    ("dixie-escalante", "Escalante", ("dixie", "escalante")),
    ("dixie-lake-powell", "Lake Powell", ("dixie", "lake powell")),
    ("dixie-pine-valley", "Pine Valley", ("dixie", "pine valley")),
    ("uwc-evanston", "Evanston", ("uwc", "evanston")),
    ("uwc-heber-kamas", "Heber/Kamas", ("uwc", "heber")),
    ("uwc-heber-kamas", "Heber/Kamas", ("uwc", "kamas")),
    ("uwc-pleasant-grove", "Pleasant Grove", ("pleasant grove",)),
    ("uwc-spanish-fork", "Spanish Fork", ("spanish fork",)),
)

BLM_DISTRICTS = (
    ("blm-cedar-city", "Cedar City", ("cedar",)),
    ("blm-grand-staircase", "Grand Staircase-Escalante", ("grand stair",)),
    ("blm-kanab", "Kanab", ("kanab",)),
)

MANTI_LA_SAL_ZONE_DISTRICTS = {
    "North Zone": (
        ("manti-la-sal-north-zone", "Manti-La Sal North Zone"),
    ),
    "South Zone - Moab and Monticello": (
        ("manti-la-sal-south-zone", "Manti-La Sal South Zone - Moab and Monticello"),
    ),
}


def clean(value: Any) -> str:
    return " ".join(str(value or "").replace("\u00a0", " ").split()).strip()


def ascii_text(value: Any) -> str:
    return unicodedata.normalize("NFKD", clean(value)).encode("ascii", "ignore").decode().lower()


def strict_identity(value: Any) -> str:
    text = ascii_text(value).replace("&", " and ")
    text = re.sub(r"\b(l\.?l\.?c\.?|inc\.?|company|co\.?)\b", " ", text)
    return re.sub(r"[^a-z0-9]+", "", text)


def slugify(value: Any) -> str:
    return re.sub(r"^-|-$", "", re.sub(r"[^a-z0-9]+", "-", ascii_text(value)))


def split_values(value: Any) -> list[str]:
    return [clean(part) for part in re.split(r"\s*[;|]\s*", clean(value)) if clean(part)]


def unique(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        key = ascii_text(value)
        if value and key not in seen:
            seen.add(key)
            result.append(value)
    return result


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_json(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        handle.write(json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def read_hunting_rows(workbook_path: Path) -> tuple[list[dict[str, str]], int]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    if "Outfitters" not in workbook.sheetnames:
        raise ValueError("Workbook is missing the Outfitters sheet")
    sheet = workbook["Outfitters"]
    headers = [clean(cell.value) for cell in sheet[1]]
    missing = sorted(REQUIRED_HEADERS - set(headers))
    if missing:
        raise ValueError(f"Workbook is missing required columns: {missing}")
    rows: list[dict[str, str]] = []
    named_count = 0
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        row = dict(zip(headers, (clean(value) for value in values)))
        if not row["OUTFITTER"]:
            continue
        named_count += 1
        if "hunting" not in ascii_text(row["HUNTING OR FISHING"]):
            continue
        if row["REVIEW STATUS"] not in REVIEW_STATUSES:
            raise ValueError(f"Unexpected review status at row {row_number}: {row['REVIEW STATUS']!r}")
        row["_source_row"] = str(row_number)
        rows.append(row)
    workbook.close()
    rows.sort(key=lambda row: ascii_text(row["OUTFITTER"]))
    return rows, named_count


def read_hunting_rows_json(path: Path) -> tuple[list[dict[str, str]], int, str, str]:
    payload = read_json(path, {})
    headers = [clean(value) for value in payload.get("headers", [])]
    missing = sorted(REQUIRED_HEADERS - set(headers))
    if missing:
        raise ValueError(f"Workbook row export is missing required columns: {missing}")
    rows: list[dict[str, str]] = []
    named_count = 0
    for row_number, values in enumerate(payload.get("rows", []), start=2):
        row = dict(zip(headers, (clean(value) for value in values)))
        if not row.get("OUTFITTER"):
            continue
        named_count += 1
        if "hunting" not in ascii_text(row.get("HUNTING OR FISHING")):
            continue
        if row.get("REVIEW STATUS") not in REVIEW_STATUSES:
            raise ValueError(f"Unexpected review status at row {row_number}: {row.get('REVIEW STATUS')!r}")
        row["_source_row"] = str(row_number)
        rows.append(row)
    rows.sort(key=lambda row: ascii_text(row["OUTFITTER"]))
    return (
        rows,
        named_count,
        clean(payload.get("sourceWorkbook")),
        clean(payload.get("sourceSha256")),
    )


def resolve_legacy_records(rows: list[dict[str, str]], legacy: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    by_exact = {ascii_text(row["OUTFITTER"]): row for row in rows}
    by_identity: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        by_identity.setdefault(strict_identity(row["OUTFITTER"]), []).append(row)

    resolved: dict[str, dict[str, Any]] = {}
    for record in legacy:
        display = clean(record.get("displayName") or record.get("businessName"))
        target = by_exact.get(ascii_text(display))
        alias = LEGACY_NAME_ALIASES.get(ascii_text(display))
        if not target and alias:
            target = by_exact.get(ascii_text(alias))
        if not target:
            candidates: dict[str, dict[str, str]] = {}
            for value in (record.get("displayName"), record.get("businessName"), record.get("legalBusinessName")):
                for candidate in by_identity.get(strict_identity(value), []):
                    candidates[candidate["OUTFITTER"]] = candidate
            if len(candidates) == 1:
                target = next(iter(candidates.values()))
        if not target:
            continue
        key = target["OUTFITTER"]
        if key in resolved:
            raise ValueError(f"Multiple legacy records resolve to {key!r}")
        resolved[key] = record
    return resolved


def normalize_usfs(raw: str) -> dict[str, Any]:
    text = ascii_text(raw)
    forest_ids: list[str] = []
    forests: list[str] = []
    district_ids: list[str] = []
    districts: list[str] = []
    for identifier, label, patterns in USFS_FORESTS:
        if any(pattern in text for pattern in patterns):
            forest_ids.append(identifier)
            forests.append(label)
    for identifier, label, patterns in USFS_DISTRICTS:
        if all(pattern in text for pattern in patterns):
            district_ids.append(identifier)
            districts.append(label)
    issues: list[str] = []
    if raw and not forest_ids:
        issues.append("unrecognized_usfs_label")
    return {
        "forestIds": unique(forest_ids),
        "forests": unique(forests),
        "districtIds": unique(district_ids),
        "districts": unique(districts),
        "issues": issues,
    }


def normalize_blm(raw: str) -> dict[str, Any]:
    text = ascii_text(raw)
    district_ids: list[str] = []
    districts: list[str] = []
    for identifier, label, patterns in BLM_DISTRICTS:
        if all(pattern in text for pattern in patterns):
            district_ids.append(identifier)
            districts.append(label)
    issues: list[str] = []
    if raw and not district_ids:
        issues.append("unrecognized_or_non_blm_label")
    return {"districtIds": unique(district_ids), "districts": unique(districts), "issues": issues}


def default_record() -> dict[str, Any]:
    return {
        "legalBusinessName": "",
        "listingType": "Outfitter",
        "publicStatus": "internal_only",
        "verificationStatus": "Unreviewed",
        "certLevel": "",
        "memberStatus": "unknown",
        "referralStatus": "hold",
        "referralPriority": "standard",
        "referralRotationGroup": "utah-internal-review",
        "branding": {"logoUrl": "", "heroImageUrl": "", "cardImageUrl": ""},
        "services": {
            "guidedHunts": True,
            "diySupport": False,
            "trespassAccess": False,
            "lodgingIncluded": False,
            "mealsIncluded": False,
            "packTrips": False,
            "airportPickup": False,
            "youthHunts": False,
            "archery": False,
            "muzzleloader": False,
            "rifle": False,
            "hamss": False,
            "otherServices": [],
        },
        "huntFit": {
            "trophyFocus": False,
            "generalSeasonFocus": False,
            "limitedEntryFocus": False,
            "oilFocus": False,
            "speciesSummary": "",
            "terrainSummary": "",
            "publicLandStrength": "",
            "accessSummary": "",
        },
        "compliance": {
            "stateLicenseNumber": "",
            "guideLicenseNumber": "",
            "outfitterLicenseNumber": "",
            "insuranceVerified": False,
            "contractOnFile": False,
            "vettingReviewedAt": "",
            "vettingReviewedBy": "",
            "nextReviewDue": "",
        },
        "publication": {
            "shortDescription": "",
            "longDescription": "",
            "whyListed": "",
            "featuredRank": 0,
            "showOnPlanner": False,
            "showOnPublicList": False,
            "showOnHomepage": False,
        },
    }


def build_master_record(
    row: dict[str, str],
    legacy_record: dict[str, Any] | None,
    workbook_name: str,
    workbook_hash: str,
    generated_at: str,
) -> tuple[dict[str, Any], dict[str, Any] | None]:
    record = copy.deepcopy(legacy_record) if legacy_record else default_record()
    identifier = clean(record.get("id")) or f"outfitter-{slugify(row['OUTFITTER'])}"
    slug = clean(record.get("slug")) or slugify(row["OUTFITTER"])
    owners = [owner.replace("*", "").strip() for owner in split_values(row["OWNER"])]
    phones = split_values(row["PHONE NUMBER"])
    emails = split_values(row["EMAIL"])
    usfs = normalize_usfs(row["SUP’s USFS"])
    blm = normalize_blm(row["SRP's BLM"])
    review_status = row["REVIEW STATUS"]
    foia_permit_confirmed = (
        "USFS FOIA 2025-FS-R4-08503-F" in row["RECONCILIATION NOTES"]
        and "issued outfitting/guiding authorization(s)" in row["RECONCILIATION NOTES"]
    )
    verification_status = clean(record.get("verificationStatus")) or "Unreviewed"
    public_eligible = review_status == "Confirmed" or verification_status == "Vetted"

    record.update(
        {
            "id": identifier,
            "slug": slug,
            "displayName": row["OUTFITTER"],
            "businessName": row["OUTFITTER"],
            "listingType": "Outfitter",
            "reviewStatus": review_status,
            "verificationStatus": verification_status,
            "publicStatus": "active" if public_eligible else "internal_only",
            "contact": {
                "primaryName": owners[0] if owners else "",
                "ownerNames": owners,
                "phonePrimary": phones[0] if phones else "",
                "phoneNumbers": phones,
                "emailPrimary": emails[0] if emails else "",
                "emailAddresses": emails,
                "website": row["WEBSITE"],
                "facebookUrl": "",
                "instagramHandle": "",
                "youtubeUrl": "",
            },
            "headquarters": {
                "city": row["LOCATION"],
                "region": "Utah",
                "state": "Utah",
                "mailingAddress": "",
                "publicMeetingLocation": "",
                "latitude": None,
                "longitude": None,
            },
            "serviceArea": {
                "speciesServed": [],
                "unitsServed": [],
                "usfsForests": usfs["forests"],
                "usfsForestIds": usfs["forestIds"],
                "usfsDistricts": usfs["districts"],
                "usfsDistrictIds": usfs["districtIds"],
                "blmDistricts": blm["districts"],
                "blmDistrictIds": blm["districtIds"],
                "countiesServed": [],
                "wmasServed": [],
                "statewide": False,
            },
        }
    )
    record.setdefault("publication", default_record()["publication"])
    record["publication"]["showOnPlanner"] = public_eligible
    record["publication"]["showOnPublicList"] = public_eligible
    if not public_eligible:
        record["publication"]["showOnHomepage"] = False
    record.setdefault("services", default_record()["services"])
    record["services"]["guidedHunts"] = True
    record.setdefault("internal", {})
    legacy_notes = [clean(note) for note in record["internal"].get("sourceNotes", []) if clean(note)]
    source_notes = unique(
        legacy_notes
        + [
            "Working master hunting-outfitter promotion",
            f"Spreadsheet review status: {review_status}",
            "Federal regions are evidence; unitsServed is intentionally blank pending direct confirmation",
        ]
    )
    legacy_repository_record_preserved = record["internal"].get("legacyRepositoryRecordPreserved")
    if not isinstance(legacy_repository_record_preserved, bool):
        legacy_repository_record_preserved = bool(legacy_record)
    record["internal"].update(
        {
            "sourceNotes": source_notes,
            "reviewNotes": unique(
                [clean(note) for note in record["internal"].get("reviewNotes", []) if clean(note)]
                + ([row["RECONCILIATION NOTES"]] if row["RECONCILIATION NOTES"] else [])
            ),
            "spreadsheetReviewStatus": review_status,
            "businessEntity": row["Business Entity"],
            "sourceWorkbook": workbook_name,
            "sourceWorkbookSha256": workbook_hash,
            "sourceSheet": "Outfitters",
            "sourceRow": int(row["_source_row"]),
            "legacyRepositoryRecordPreserved": legacy_repository_record_preserved,
            "lastNormalizedAt": generated_at[:10],
            "lastEditedBy": "Codex",
            "federalPermitEvidenceStatus": "Confirmed" if foia_permit_confirmed else "Needs Verification",
        }
    )

    evidence = None
    if row["SUP’s USFS"] or row["SRP's BLM"]:
        issues = unique(usfs["issues"] + blm["issues"])
        evidence = {
            "outfitterId": identifier,
            "displayName": row["OUTFITTER"],
            "reviewStatus": review_status,
            "evidenceBasis": (
                "USDA_FOREST_SERVICE_ISSUED_AUTHORIZATION_ROSTER"
                if foia_permit_confirmed
                else "SPREADSHEET_REPORTED_FEDERAL_PERMIT_AREA"
            ),
            "permitEvidenceStatus": "Confirmed" if foia_permit_confirmed else "Needs Verification",
            "relationship": "FEDERAL_SERVICE_AREA_EVIDENCE",
            "unitServiceClaimStatus": "NOT_CONFIRMED",
            "usfsRaw": row["SUP’s USFS"],
            "usfsForestIds": usfs["forestIds"],
            "usfsForests": usfs["forests"],
            "usfsDistrictIds": usfs["districtIds"],
            "usfsDistricts": usfs["districts"],
            "confirmedUsfsForestIds": usfs["forestIds"] if foia_permit_confirmed else [],
            "confirmedUsfsDistrictIds": usfs["districtIds"] if foia_permit_confirmed else [],
            "blmRaw": row["SRP's BLM"],
            "blmDistrictIds": blm["districtIds"],
            "blmDistricts": blm["districts"],
            "confirmedBlmDistrictIds": [],
            "normalizationStatus": "NEEDS_VERIFICATION" if issues else "NORMALIZED",
            "normalizationIssues": issues,
            "sourceWorkbookSha256": workbook_hash,
            "sourceRow": int(row["_source_row"]),
        }
    return record, evidence


def apply_authoritative_federal_evidence(
    master_rows: list[dict[str, Any]],
    evidence_rows: list[dict[str, Any]],
) -> tuple[int, int]:
    authority = read_json(MANTI_LA_SAL_AUTHORITY_PATH, {})
    metadata = authority.get("metadata", {})
    authority_records = authority.get("records", [])
    if len(authority_records) != 19:
        raise ValueError(f"Expected 19 Manti-La Sal authority matches, found {len(authority_records)}")

    master_by_id = {row["id"]: row for row in master_rows}
    evidence_by_id = {row["outfitterId"]: row for row in evidence_rows}
    spreadsheet_evidence_count = len(evidence_rows)
    for authority_row in authority_records:
        outfitter_id = authority_row["outfitterId"]
        record = master_by_id.get(outfitter_id)
        if not record:
            raise ValueError(f"Manti-La Sal authority record references missing outfitter: {outfitter_id}")
        if record["displayName"] != authority_row["masterDisplayName"]:
            raise ValueError(f"Manti-La Sal authority display-name mismatch for {outfitter_id}")

        service_area = record["serviceArea"]
        zone_districts = MANTI_LA_SAL_ZONE_DISTRICTS.get(authority_row["zone"])
        if not zone_districts:
            raise ValueError(f"Unmapped Manti-La Sal permit zone: {authority_row['zone']}")
        service_area["usfsForestIds"] = [
            value for value in service_area["usfsForestIds"] if value != "manti-la-sal"
        ]
        service_area["usfsForests"] = [
            value for value in service_area["usfsForests"] if value != "Manti-La Sal"
        ]
        service_area["usfsDistrictIds"] = unique(
            service_area["usfsDistrictIds"] + [identifier for identifier, _ in zone_districts]
        )
        service_area["usfsDistricts"] = unique(
            service_area["usfsDistricts"] + [label for _, label in zone_districts]
        )
        official_evidence = {
            "sourceId": metadata["id"],
            "authority": metadata["authority"],
            "sourceDocument": metadata["sourceDocument"],
            "sourceDocumentSha256": metadata["sourceDocumentSha256"],
            "sourceLastUpdated": metadata["sourceLastUpdated"],
            "officialPermitHolderName": authority_row["officialPermitHolderName"],
            "zone": authority_row["zone"],
            "permitScope": authority_row["permitScope"],
            "claim": "CURRENT_FOREST_SERVICE_PERMIT_HOLDER",
        }
        record["internal"]["authoritativeFederalPermitEvidence"] = [official_evidence]
        record["internal"]["federalPermitEvidenceStatus"] = "Confirmed"
        record["internal"]["sourceNotes"] = unique(
            record["internal"]["sourceNotes"]
            + [
                "USDA Forest Service confirms current Manti-La Sal permit-holder status; "
                f"source last updated {metadata['sourceLastUpdated']}"
            ]
        )

        evidence = evidence_by_id.get(outfitter_id)
        if evidence is None:
            evidence = {
                "outfitterId": outfitter_id,
                "displayName": record["displayName"],
                "reviewStatus": record["reviewStatus"],
                "evidenceBasis": "USDA_FOREST_SERVICE_CURRENT_PERMIT_HOLDER",
                "permitEvidenceStatus": "Confirmed",
                "relationship": "FEDERAL_SERVICE_AREA_EVIDENCE",
                "unitServiceClaimStatus": "NOT_CONFIRMED",
                "usfsRaw": "",
                "usfsForestIds": [],
                "usfsForests": [],
                "usfsDistrictIds": [],
                "usfsDistricts": [],
                "confirmedUsfsForestIds": [],
                "confirmedUsfsDistrictIds": [],
                "blmRaw": "",
                "blmDistrictIds": [],
                "blmDistricts": [],
                "confirmedBlmDistrictIds": [],
                "normalizationStatus": "NORMALIZED",
                "normalizationIssues": [],
                "sourceWorkbookSha256": record["internal"]["sourceWorkbookSha256"],
                "sourceRow": record["internal"]["sourceRow"],
            }
            evidence_rows.append(evidence)
            evidence_by_id[outfitter_id] = evidence
        else:
            evidence["evidenceBasis"] = "MULTIPLE_SOURCES"
        evidence["permitEvidenceStatus"] = "Confirmed"
        evidence["usfsForestIds"] = [
            value for value in evidence["usfsForestIds"] if value != "manti-la-sal"
        ]
        evidence["usfsForests"] = [
            value for value in evidence["usfsForests"] if value != "Manti-La Sal"
        ]
        evidence["usfsDistrictIds"] = unique(
            evidence["usfsDistrictIds"] + [identifier for identifier, _ in zone_districts]
        )
        evidence["usfsDistricts"] = unique(
            evidence["usfsDistricts"] + [label for _, label in zone_districts]
        )
        evidence["confirmedUsfsDistrictIds"] = unique(
            evidence.get("confirmedUsfsDistrictIds", [])
            + [identifier for identifier, _ in zone_districts]
        )
        evidence.setdefault("evidenceSources", [])
        if evidence.get("usfsRaw") or evidence.get("blmRaw"):
            workbook_source = {
                "sourceId": "reconciled-working-master",
                "basis": "SPREADSHEET_REPORTED_FEDERAL_PERMIT_AREA",
                "sourceWorkbookSha256": evidence["sourceWorkbookSha256"],
                "sourceRow": evidence["sourceRow"],
            }
            if workbook_source not in evidence["evidenceSources"]:
                evidence["evidenceSources"].append(workbook_source)
        evidence["evidenceSources"].append(official_evidence)

    evidence_rows.sort(key=lambda item: ascii_text(item["displayName"]))
    return spreadsheet_evidence_count, len(authority_records)


def apply_fishlake_authoritative_evidence(
    master_rows: list[dict[str, Any]],
    evidence_rows: list[dict[str, Any]],
) -> int:
    authority = read_json(FISHLAKE_AUTHORITY_PATH, {})
    metadata = authority.get("metadata", {})
    authority_records = authority.get("records", [])
    if len(authority_records) != 41:
        raise ValueError(f"Expected 41 Fishlake authority matches, found {len(authority_records)}")

    master_by_id = {row["id"]: row for row in master_rows}
    evidence_by_id = {row["outfitterId"]: row for row in evidence_rows}
    for authority_row in authority_records:
        outfitter_id = authority_row["outfitterId"]
        record = master_by_id.get(outfitter_id)
        if not record:
            raise ValueError(f"Fishlake authority record references missing outfitter: {outfitter_id}")
        if record["displayName"] != authority_row["masterDisplayName"]:
            raise ValueError(f"Fishlake authority display-name mismatch for {outfitter_id}")

        service_area = record["serviceArea"]
        service_area["usfsForestIds"] = unique(service_area["usfsForestIds"] + ["fishlake"])
        service_area["usfsForests"] = unique(service_area["usfsForests"] + ["Fishlake"])
        official_evidence = {
            "sourceId": metadata["id"],
            "authority": metadata["authority"],
            "sourceDocument": metadata["sourceDocument"],
            "sourceDocumentSha256": metadata["sourceDocumentSha256"],
            "sourceAsOf": metadata["sourceAsOf"],
            "officialPermitHolderName": authority_row["officialPermitHolderName"],
            "matchType": authority_row["matchType"],
            "permitScope": metadata["permitScope"],
            "claim": "CURRENT_FOREST_SERVICE_PERMIT_HOLDER",
        }
        record["internal"].setdefault("authoritativeFederalPermitEvidence", []).append(
            official_evidence
        )
        record["internal"]["federalPermitEvidenceStatus"] = "Confirmed"
        record["internal"]["sourceNotes"] = unique(
            record["internal"]["sourceNotes"]
            + [
                "USDA Forest Service Fishlake roster confirms permit-holder status; "
                f"source identified as {metadata['sourceAsOf']}"
            ]
        )

        evidence = evidence_by_id.get(outfitter_id)
        if evidence is None:
            evidence = {
                "outfitterId": outfitter_id,
                "displayName": record["displayName"],
                "reviewStatus": record["reviewStatus"],
                "evidenceBasis": "USDA_FOREST_SERVICE_CURRENT_PERMIT_HOLDER",
                "permitEvidenceStatus": "Confirmed",
                "relationship": "FEDERAL_SERVICE_AREA_EVIDENCE",
                "unitServiceClaimStatus": "NOT_CONFIRMED",
                "usfsRaw": "",
                "usfsForestIds": [],
                "usfsForests": [],
                "usfsDistrictIds": [],
                "usfsDistricts": [],
                "confirmedUsfsForestIds": [],
                "confirmedUsfsDistrictIds": [],
                "blmRaw": "",
                "blmDistrictIds": [],
                "blmDistricts": [],
                "confirmedBlmDistrictIds": [],
                "normalizationStatus": "NORMALIZED",
                "normalizationIssues": [],
                "sourceWorkbookSha256": record["internal"]["sourceWorkbookSha256"],
                "sourceRow": record["internal"]["sourceRow"],
            }
            evidence_rows.append(evidence)
            evidence_by_id[outfitter_id] = evidence
        else:
            evidence["evidenceBasis"] = "MULTIPLE_SOURCES"
        evidence["permitEvidenceStatus"] = "Confirmed"
        evidence["usfsForestIds"] = unique(evidence["usfsForestIds"] + ["fishlake"])
        evidence["usfsForests"] = unique(evidence["usfsForests"] + ["Fishlake"])
        evidence["confirmedUsfsForestIds"] = unique(
            evidence.get("confirmedUsfsForestIds", []) + ["fishlake"]
        )
        evidence.setdefault("evidenceSources", [])
        if evidence.get("usfsRaw") or evidence.get("blmRaw"):
            workbook_source = {
                "sourceId": "reconciled-working-master",
                "basis": "SPREADSHEET_REPORTED_FEDERAL_PERMIT_AREA",
                "sourceWorkbookSha256": evidence["sourceWorkbookSha256"],
                "sourceRow": evidence["sourceRow"],
            }
            if workbook_source not in evidence["evidenceSources"]:
                evidence["evidenceSources"].append(workbook_source)
        evidence["evidenceSources"].append(official_evidence)

    evidence_rows.sort(key=lambda item: ascii_text(item["displayName"]))
    return len(authority_records)


def apply_operator_confirmed_service_claims(
    master_rows: list[dict[str, Any]], evidence_rows: list[dict[str, Any]]
) -> int:
    source = read_json(WILD_EYEZ_CONFIRMED_SERVICE_PATH, {})
    outfitter_id = source["outfitterId"]
    master = next((row for row in master_rows if row["id"] == outfitter_id), None)
    evidence = next((row for row in evidence_rows if row["outfitterId"] == outfitter_id), None)
    if not master or not evidence:
        raise ValueError("Wild Eyez operator-confirmed source does not resolve to the internal master")
    if master["displayName"] != source["displayName"]:
        raise ValueError("Wild Eyez operator-confirmed display name does not match the internal master")

    scope = source["federalAuthorizationScope"]
    service_area = master["serviceArea"]
    service_area["usfsForestIds"] = unique(
        [value for value in service_area["usfsForestIds"] if value not in {"manti-la-sal", "uwc"}]
        + scope["wholeForestIds"]
    )
    service_area["usfsDistrictIds"] = unique(
        service_area["usfsDistrictIds"] + scope["rangerDistrictIds"]
    )
    service_area["usfsDistricts"] = unique(
        service_area["usfsDistricts"] + scope["rangerDistricts"]
    )
    master["internal"]["confirmedUnitServiceClaims"] = source["confirmedServiceClaims"]
    master["internal"]["federalPermitEvidenceStatus"] = "Confirmed"
    master["internal"]["confirmedServiceClaimSource"] = str(
        WILD_EYEZ_CONFIRMED_SERVICE_PATH.relative_to(ROOT)
    ).replace("\\", "/")

    evidence["usfsForestIds"] = unique(
        [value for value in evidence["usfsForestIds"] if value not in {"manti-la-sal", "uwc"}]
        + scope["wholeForestIds"]
    )
    evidence["usfsDistrictIds"] = unique(
        evidence["usfsDistrictIds"] + scope["rangerDistrictIds"]
    )
    evidence["usfsDistricts"] = unique(evidence["usfsDistricts"] + scope["rangerDistricts"])
    evidence["confirmedUnitServiceClaims"] = source["confirmedServiceClaims"]
    evidence["permitEvidenceStatus"] = "Confirmed"
    evidence["confirmedUsfsForestIds"] = unique(
        evidence.get("confirmedUsfsForestIds", []) + scope["wholeForestIds"]
    )
    evidence["confirmedUsfsDistrictIds"] = unique(
        evidence.get("confirmedUsfsDistrictIds", []) + scope["rangerDistrictIds"]
    )
    evidence.setdefault("evidenceSources", []).append(
        {
            "sourceId": source["metadata"]["id"],
            "authority": source["metadata"]["authority"],
            "confirmedAt": source["metadata"]["confirmedAt"],
            "relationship": source["metadata"]["relationship"],
            "sourceFile": str(WILD_EYEZ_CONFIRMED_SERVICE_PATH.relative_to(ROOT)).replace("\\", "/"),
        }
    )
    return len(source["confirmedServiceClaims"])


def build_public_repository_master(
    public_rows: list[dict[str, Any]], internal_rows: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in public_rows:
        grouped.setdefault(clean(row.get("listingName")), []).append(row)
    internal_by_id = {row["id"]: row for row in internal_rows}
    result: list[dict[str, Any]] = []
    for listing_name, rows in grouped.items():
        identifier = f"outfitter-{slugify(listing_name)}"
        internal = internal_by_id.get(identifier, {})
        record = default_record()
        record.update(
            {
                "id": identifier,
                "slug": slugify(listing_name),
                "displayName": listing_name,
                "businessName": listing_name,
                "legalBusinessName": "",
                "listingType": "Outfitter",
                "publicStatus": "active",
                "verificationStatus": "Vetted",
                "reviewStatus": clean(internal.get("reviewStatus")),
                "certLevel": next((clean(row.get("certLevel")) for row in rows if clean(row.get("certLevel"))), ""),
                "referralStatus": "eligible",
                "contact": {
                    "primaryName": next(
                        (clean(name) for row in rows for name in row.get("ownerName", []) if clean(name)), ""
                    ),
                    "ownerNames": unique(
                        [clean(name) for row in rows for name in row.get("ownerName", []) if clean(name)]
                    ),
                    "phonePrimary": next(
                        (clean(phone) for row in rows for phone in row.get("phone", []) if clean(phone)), ""
                    ),
                    "phoneNumbers": unique(
                        [clean(phone) for row in rows for phone in row.get("phone", []) if clean(phone)]
                    ),
                    "emailPrimary": next(
                        (clean(email) for row in rows for email in row.get("email", []) if clean(email)), ""
                    ),
                    "emailAddresses": unique(
                        [clean(email) for row in rows for email in row.get("email", []) if clean(email)]
                    ),
                    "website": next((clean(row.get("website")) for row in rows if clean(row.get("website"))), ""),
                    "facebookUrl": "",
                    "instagramHandle": "",
                    "youtubeUrl": "",
                },
                "branding": {
                    "logoUrl": next((clean(row.get("logoUrl")) for row in rows if clean(row.get("logoUrl"))), ""),
                    "heroImageUrl": "",
                    "cardImageUrl": "",
                },
                "headquarters": {
                    "city": next((clean(row.get("city")) for row in rows if clean(row.get("city"))), ""),
                    "region": next((clean(row.get("region")) for row in rows if clean(row.get("region"))), "Utah"),
                    "state": "Utah",
                    "mailingAddress": "",
                    "publicMeetingLocation": "",
                    "latitude": None,
                    "longitude": None,
                },
                "serviceArea": {
                    "speciesServed": unique(
                        [clean(value) for row in rows for value in row.get("speciesServed", []) if clean(value)]
                    ),
                    "unitsServed": unique(
                        [clean(value) for row in rows for value in row.get("unitsServed", []) if clean(value)]
                    ),
                    "usfsForests": unique(
                        [clean(value) for row in rows for value in row.get("usfsForests", []) if clean(value)]
                    ),
                    "blmDistricts": unique(
                        [clean(value) for row in rows for value in row.get("blmDistricts", []) if clean(value)]
                    ),
                    "countiesServed": [],
                    "wmasServed": [],
                    "statewide": False,
                },
            }
        )
        record["services"]["guidedHunts"] = True
        record["publication"].update(
            {"showOnPlanner": True, "showOnPublicList": True, "showOnHomepage": False}
        )
        result.append(record)
    return sorted(result, key=lambda item: ascii_text(item["displayName"]))


def split_ids(value: Any) -> set[str]:
    return {clean(item).lower() for item in clean(value).split("|") if clean(item)}


def coverage_base() -> list[dict[str, Any]]:
    rows = read_json(COVERAGE_PATH, [])
    if rows:
        return rows
    canonical = read_json(CANONICAL_PATH, {})
    return canonical.get("outfitters", {}).get("federal_coverage", [])


def rebuild_coverage(base_rows: list[dict[str, Any]], evidence_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rebuilt: list[dict[str, Any]] = []
    for source_row in base_rows:
        row = copy.deepcopy(source_row)
        unit_usfs = split_ids(row.get("PrimaryUsfsForestId"))
        unit_blm = split_ids(row.get("PrimaryBlmDistrictId"))
        usfs_matches = [item for item in evidence_rows if unit_usfs & set(item["usfsForestIds"])]
        blm_matches = [item for item in evidence_rows if unit_blm & set(item["blmDistrictIds"])]
        federal_by_id = {item["outfitterId"]: item for item in usfs_matches + blm_matches}
        federal_matches = sorted(federal_by_id.values(), key=lambda item: ascii_text(item["displayName"]))
        usfs_matches.sort(key=lambda item: ascii_text(item["displayName"]))
        blm_matches.sort(key=lambda item: ascii_text(item["displayName"]))
        row.update(
            {
                "UsfsPermitMatchedOutfitterCount": len(usfs_matches),
                "UsfsPermitMatchedOutfitters": " | ".join(item["displayName"] for item in usfs_matches),
                "UsfsPermitMatchedOutfitterIds": [item["outfitterId"] for item in usfs_matches],
                "BlmPermitMatchedOutfitterCount": len(blm_matches),
                "BlmPermitMatchedOutfitters": " | ".join(item["displayName"] for item in blm_matches),
                "BlmPermitMatchedOutfitterIds": [item["outfitterId"] for item in blm_matches],
                "FederalPermitMatchedOutfitterCount": len(federal_matches),
                "FederalPermitMatchedOutfitters": " | ".join(item["displayName"] for item in federal_matches),
                "FederalPermitMatchedOutfitterIds": [item["outfitterId"] for item in federal_matches],
                "CoverageRelationship": "FEDERAL_PERMIT_OVERLAP_ONLY",
                "UnitServiceClaimStatus": "NOT_CONFIRMED",
                "CoverageEvidenceSource": "local_data/outfitters/outfitter-federal-service-area-evidence.internal.json",
                "Notes": (
                    "Provisional federal-permit overlap only. A USFS forest or BLM district overlap does not confirm "
                    "that an outfitter serves the complete DWR unit. Direct outfitter, permit, or authoritative "
                    "geographic evidence is required before populating unitsServed."
                ),
            }
        )
        rebuilt.append(row)
    return rebuilt


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--workbook-rows-json", type=Path)
    parser.add_argument("--expected-sha256", default="")
    parser.add_argument("--generated-at", default="")
    args = parser.parse_args()

    if bool(args.workbook) == bool(args.workbook_rows_json):
        raise ValueError("Specify exactly one of --workbook or --workbook-rows-json")
    if args.workbook_rows_json:
        rows, named_count, workbook_name, workbook_hash = read_hunting_rows_json(
            args.workbook_rows_json.resolve()
        )
        workbook_path = Path(workbook_name)
    else:
        workbook_path = args.workbook.resolve()
        workbook_hash = sha256(workbook_path)
        rows, named_count = read_hunting_rows(workbook_path)
        workbook_name = workbook_path.name
    if args.expected_sha256 and workbook_hash.lower() != args.expected_sha256.lower():
        raise ValueError(f"Workbook SHA-256 mismatch: {workbook_hash}")
    generated_at = args.generated_at or datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")

    legacy = read_json(MASTER_PATH, read_json(LEGACY_MASTER_PATH, []))
    legacy_by_name = resolve_legacy_records(rows, legacy)
    master_rows: list[dict[str, Any]] = []
    evidence_rows: list[dict[str, Any]] = []
    for row in rows:
        record, evidence = build_master_record(
            row,
            legacy_by_name.get(row["OUTFITTER"]),
            workbook_name,
            workbook_hash,
            generated_at,
        )
        master_rows.append(record)
        if evidence:
            evidence_rows.append(evidence)

    spreadsheet_evidence_count, manti_authority_evidence_count = apply_authoritative_federal_evidence(
        master_rows, evidence_rows
    )
    fishlake_authority_evidence_count = apply_fishlake_authoritative_evidence(
        master_rows, evidence_rows
    )
    authority_evidence_count = (
        manti_authority_evidence_count + fishlake_authority_evidence_count
    )
    operator_confirmed_service_claim_count = apply_operator_confirmed_service_claims(
        master_rows, evidence_rows
    )

    if len(master_rows) != 138:
        raise ValueError(f"Expected 138 named hunting businesses, found {len(master_rows)}")
    if named_count != 182:
        raise ValueError(f"Expected 182 named workbook businesses, found {named_count}")
    if spreadsheet_evidence_count < 1 or len(evidence_rows) < spreadsheet_evidence_count:
        raise ValueError("Federal service-area evidence counts are internally inconsistent")
    if len({row["id"] for row in master_rows}) != len(master_rows):
        raise ValueError("Duplicate outfitter ids produced")
    if any(row["serviceArea"]["unitsServed"] for row in master_rows):
        raise ValueError("unitsServed must remain empty in the promoted internal master")

    evidence_payload = {
        "metadata": {
            "id": "outfitter-federal-service-area-evidence",
            "generatedAt": generated_at,
            "sourceWorkbook": workbook_name,
            "sourceWorkbookSha256": workbook_hash,
            "sourceSheet": "Outfitters",
            "sourceDocuments": [
                {
                    "sourceId": "manti-la-sal-permitted-hunting-outfitters-2026",
                    "authority": "USDA Forest Service",
                    "sourceFile": "data/source-evidence/manti-la-sal-permitted-hunting-outfitters-2026.json",
                    "sourceDocument": "Manti-La Sal National Forest Permits Forest Service.pdf",
                    "sourceDocumentSha256": "65e2865f7c6ed744908e39301cda536032568e0f39d8e7505898367fd0f244e7",
                    "sourceLastUpdated": "2026-05-19",
                },
                {
                    "sourceId": "fishlake-permitted-hunting-outfitters-2025",
                    "authority": "USDA Forest Service",
                    "sourceFile": "data/source-evidence/fishlake-permitted-hunting-outfitters-2025.json",
                    "sourceDocument": "Fishlake Outfitters 2025.JPG",
                    "sourceDocumentSha256": "605b1ba01bc1aa24b666106cf4410dbbc0057c77d4143a7ae821d4a998f521ad",
                    "sourceAsOf": "2025",
                    "permitScope": "Fishlake National Forest - All Forest",
                },
                {
                    "sourceId": "wild-eyez-confirmed-elk-service-area-2026",
                    "authority": "Wild Eyez Outfitters operator confirmation",
                    "sourceFile": "data/source-evidence/wild-eyez-confirmed-elk-service-area-2026.json",
                    "confirmedAt": "2026-09-08"
                }
            ],
            "namedWorkbookBusinesses": named_count,
            "promotedHuntingBusinesses": len(master_rows),
            "spreadsheetFederalEvidenceRecords": spreadsheet_evidence_count,
            "authoritativePermitEvidenceRecords": authority_evidence_count,
            "operatorConfirmedServiceClaims": operator_confirmed_service_claim_count,
            "federalEvidenceRecords": len(evidence_rows),
            "excludedRule": "Fishing-only and unclassified businesses are not promoted",
            "unitAssociationBoundary": "Federal service-area evidence is not a confirmed DWR unit service claim",
        },
        "records": evidence_rows,
    }
    rebuilt_coverage = rebuild_coverage(coverage_base(), evidence_rows)
    if len(rebuilt_coverage) != 761:
        raise ValueError(f"Expected 761 federal coverage rows, found {len(rebuilt_coverage)}")

    write_json(MASTER_PATH, master_rows)
    write_json(EVIDENCE_PATH, evidence_payload)
    write_json(COVERAGE_PATH, rebuilt_coverage)

    public_rows = read_json(PUBLIC_PATH, [])
    vetted_public_rows = [
        row for row in public_rows if row.get("verificationStatus") == "Vetted"
    ]
    if len(vetted_public_rows) != 11:
        raise ValueError("Existing public feed must retain the 11-row vetted contact contract")
    versioned_master = build_public_repository_master(vetted_public_rows, master_rows)
    write_json(LEGACY_MASTER_PATH, versioned_master)
    if any(
        row.get("verificationStatus") != "Vetted" or not row.get("services", {}).get("guidedHunts")
        for row in versioned_master
    ):
        raise ValueError("The versioned public-repository master must contain only vetted hunting records")

    write_json(
        MANIFEST_PATH,
        {
            "metadata": {
                "id": "outfitter-internal-master-manifest",
                "generatedAt": generated_at,
                "privacyBoundary": "The full internal records are local and Git-ignored because this repository is public",
            },
            "source": {
                "workbook": workbook_name,
                "workbookSha256": workbook_hash,
                "sheet": "Outfitters",
            },
            "localArtifacts": {
                "master": "local_data/outfitters/outfitters-master.internal.json",
                "federalEvidence": "local_data/outfitters/outfitter-federal-service-area-evidence.internal.json",
                "provisionalCoverage": "local_data/outfitters/outfitter-federal-unit-coverage-review.internal.json",
            },
            "counts": {
                "namedWorkbookBusinesses": named_count,
                "promotedHuntingBusinesses": len(master_rows),
                "reviewStatus": {
                    status: sum(row["reviewStatus"] == status for row in master_rows)
                    for status in sorted(REVIEW_STATUSES)
                },
                "spreadsheetFederalEvidenceRecords": spreadsheet_evidence_count,
                "authoritativePermitEvidenceRecords": authority_evidence_count,
                "operatorConfirmedServiceClaims": operator_confirmed_service_claim_count,
                "uniqueFederalEvidenceRecords": len(evidence_rows),
                "provisionalCoverageRows": len(rebuilt_coverage),
                "publicRowsUnchanged": len(vetted_public_rows),
                "publicRepositoryMasterBusinesses": len(versioned_master),
            },
            "rules": {
                "excluded": "Fishing-only and unclassified businesses",
                "unitAssociation": "FEDERAL_PERMIT_OVERLAP_ONLY; NOT_CONFIRMED",
                "publicSource": "data/outfitters-public.json",
            },
        },
    )

    print(
        json.dumps(
            {
                "ok": True,
                "source_sha256": workbook_hash,
                "named_workbook_businesses": named_count,
                "promoted_hunting_businesses": len(master_rows),
                "review_status_counts": {
                    status: sum(row["reviewStatus"] == status for row in master_rows)
                    for status in sorted(REVIEW_STATUSES)
                },
                "legacy_records_safely_preserved": sum(
                    bool(row["internal"]["legacyRepositoryRecordPreserved"]) for row in master_rows
                ),
                "spreadsheet_federal_evidence_records": spreadsheet_evidence_count,
                "authoritative_permit_evidence_records": authority_evidence_count,
                "operator_confirmed_service_claims": operator_confirmed_service_claim_count,
                "unique_federal_evidence_records": len(evidence_rows),
                "federal_overlap_rows": len(rebuilt_coverage),
                "federal_overlap_rows_with_matches": sum(
                    int(row["FederalPermitMatchedOutfitterCount"]) > 0 for row in rebuilt_coverage
                ),
                "public_rows_unchanged": len(vetted_public_rows),
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
